"""Download official PDFs listed in the team source-register workbook."""

from __future__ import annotations

import argparse
import csv
import json
import shutil
import urllib.error
import urllib.request
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


REGISTER_FIELDS = [
    "doc_id", "title", "issuer", "authority_level", "doc_type", "version",
    "effective_from", "effective_to", "supersedes", "scope_programmes",
    "scope_batches", "provenance", "retrieved_on", "synthetic",
]

# Five links in the supplied workbook were visibly truncated by the spreadsheet
# export. These full links were recovered from NSUT's official notices archive.
URL_OVERRIDES = {
    "NSUT-20260806-CORRIEV": "https://www.imsnsit.org/imsnsit/plum_url.php?0HZDWp/1cALdt+O9qpddY6cIdMn/GEWZMbcMRYOaYrLBKqiHVgZDEthR9v+1jxbj946Ox0BWN+0sGKZWQBjjZAtA4yTJFNVHEyTczhk7z13cAKIUfWDm6dKGsoRIMZgLR/rhVPT1T809hODFRNZoxAVR/C524YH5nD8VgbpR6EjS0aRypzmnIWRctkP6bLH59m6CfiLvWPVd2D6jEWL3oSgiVHBJljAn9NEVaaAcoBo=",
    "NSUT-20260924-CORRIEV23": "https://www.imsnsit.org/imsnsit/plum_url.php?a6mu91rMoJDloC3iyXxMXFhzZj/sMFb2anut+mivm56u8J3FqD8qqwPpEwynKRlT/T0JJYPYL+LFHo/VOWL9fHDghBXLcC40WT7AXtwEOLXUtwFqjSl0DoQjs3rczC7g6xU/jWlOz0Xq5QXFSbne8HKrN/naj0Uvp9404aWPQat18dPnL9wmQbR790hfJ+GE0uXYFPnFaa0lY0NipFuGtBYaz8mmRVpUFHsIxxYM3po=",
    "NSUT-20260602-NOINSTAL": "https://www.imsnsit.org/imsnsit/plum_url.php?w6e6EsuItYkLPKRZANWe2fKeipUPbhD55XLuT4SWGbT58wFlyPsZtyc6478l9i9WoF4HqEKMcR4omLesS/Zx+lA3wno9b62El0Kr60ibo0SyNdq9DUwWdWG1nx893d+h+1xW4lN0762+EuU1anw/Ngy432zdG/bPWHMRHBSnjVUzE3OW5ysgfu6OTyUBbIe+Z+jzhlQeNd6wy/RkV0Avjw==",
    "NSUT-20260611-NSP": "https://www.imsnsit.org/imsnsit/plum_url.php?ZbUdfXmfbwYfVF+IU3jM9TkEeMQ33b+cOiavB4n+tW642C8hPjWU6kBHaP0gx0iwUirv0PKeg1ah9q/3Dv0vZJYxZrj0P1lJztfGcH05wnQk9ZPOIaluHQaSWD9/rG4+FvYqoXqAgVg0W8AwZ9VJkEnM642kihztUfCOqvK21PMPe2dMP5i5HQeoIZ7/q/dwq5aqctyeNT4KrUKMl7Q47Q2uy99z7AgTmM3QXTrqzaBsMc8hAsDg6l8vvA5g5FAp",
    "NSUT-20261001-EDIST": "https://www.imsnsit.org/imsnsit/plum_url.php?5Cvv3fS25/yzgFdGSyk6kNvoGmw0D6zkE7Wdehj5uHYlcKb6988ZEb2uXNtP0hgqWpNSnIIs+y7F1h/LijAs3U2hOLyJm27mpdvkJDAeOJV4hDUfAg+pWJDWA5+8vrjudTxois8r2QDt7Me5wJo5YOm0dNMFuDsJvN3+cgSWanFza2bZGqh6Qo43AuCdvkqn",
}


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def read_register(workbook_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    rows = sheet.iter_rows(values_only=True)
    headers = [_text(value) for value in next(rows)]
    missing = [field for field in REGISTER_FIELDS if field not in headers]
    if missing:
        raise ValueError(f"Source workbook is missing columns: {', '.join(missing)}")

    output = []
    for values in rows:
        raw = dict(zip(headers, values))
        row = {field: _text(raw.get(field)) for field in REGISTER_FIELDS}
        if row["doc_id"]:
            if row["doc_id"] in URL_OVERRIDES:
                row["provenance"] = URL_OVERRIDES[row["doc_id"]]
            output.append(row)
    workbook.close()
    return output


def download_pdf(url: str, destination: Path) -> dict:
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0",
            "Accept": "application/pdf,*/*;q=0.8",
            "Referer": "https://www.imsnsit.org/imsnsit/notifications.php",
        },
    )
    temporary = destination.with_suffix(".pdf.part")
    try:
        with urllib.request.urlopen(request, timeout=60) as response, temporary.open("wb") as output:
            shutil.copyfileobj(response, output)
            content_type = response.headers.get("Content-Type", "")
            final_url = response.geturl()
        with temporary.open("rb") as source:
            signature = source.read(5)
        if signature != b"%PDF-":
            temporary.unlink(missing_ok=True)
            raise ValueError(f"response is not a PDF (Content-Type: {content_type})")
        temporary.replace(destination)
        return {
            "status": "downloaded",
            "bytes": destination.stat().st_size,
            "content_type": content_type,
            "final_url": final_url,
        }
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path("data/documents"))
    parser.add_argument("--register", type=Path, default=Path("data/source_register.csv"))
    parser.add_argument(
        "--manifest", type=Path, default=Path("data/source_download_manifest.json")
    )
    args = parser.parse_args()

    rows = read_register(args.workbook)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.register.parent.mkdir(parents=True, exist_ok=True)
    with args.register.open("w", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(output, fieldnames=REGISTER_FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    manifest = []
    for row in rows:
        destination = args.output_dir / f"{row['doc_id']}.pdf"
        item = {"doc_id": row["doc_id"], "path": str(destination)}
        try:
            item.update(download_pdf(row["provenance"], destination))
            print(f"downloaded {row['doc_id']} ({item['bytes']} bytes)")
        except (urllib.error.URLError, TimeoutError, ValueError, OSError) as exc:
            item.update(status="failed", error=f"{type(exc).__name__}: {exc}")
            print(f"FAILED {row['doc_id']}: {exc}")
        manifest.append(item)

    args.manifest.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    succeeded = sum(item["status"] == "downloaded" for item in manifest)
    print(f"downloaded {succeeded}/{len(manifest)} PDFs")
    return 0 if succeeded == len(manifest) else 1


if __name__ == "__main__":
    raise SystemExit(main())
